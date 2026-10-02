"""P3-5: app/services/document_processor.py 覆盖率补充测试。"""
import os

import pytest

from app.config import settings
from app.services.document_processor import process_document


@pytest.fixture
def upload_root(tmp_path, monkeypatch):
    """将 UPLOAD_ROOT 指向临时目录。"""
    monkeypatch.setattr(settings, "UPLOAD_ROOT", str(tmp_path))
    return tmp_path


class TestProcessTextFiles:
    """文本类文件解析测试。"""

    @pytest.mark.asyncio
    async def test_process_txt(self, upload_root):
        f = upload_root / "test.txt"
        f.write_text("Hello, World!", encoding="utf-8")
        result = await process_document(str(f))
        assert "Hello, World!" in result

    @pytest.mark.asyncio
    async def test_process_csv(self, upload_root):
        f = upload_root / "test.csv"
        f.write_text("a,b,c\n1,2,3", encoding="utf-8")
        result = await process_document(str(f))
        assert "a,b,c" in result

    @pytest.mark.asyncio
    async def test_process_json(self, upload_root):
        f = upload_root / "test.json"
        f.write_text('{"key": "value"}', encoding="utf-8")
        result = await process_document(str(f))
        assert '"key": "value"' in result


class TestProcessPdf:
    """PDF 解析测试。"""

    @pytest.mark.asyncio
    async def test_process_pdf(self, upload_root):
        try:
            from pypdf import PdfWriter
        except ImportError:
            pytest.skip("pypdf 未安装")

        f = upload_root / "test.pdf"
        writer = PdfWriter()
        writer.add_blank_page(width=72, height=72)
        with open(f, "wb") as fh:
            writer.write(fh)

        result = await process_document(str(f))
        assert isinstance(result, str)


class TestProcessDocx:
    """DOCX 解析测试。"""

    @pytest.mark.asyncio
    async def test_process_docx(self, upload_root):
        try:
            from docx import Document
        except ImportError:
            pytest.skip("python-docx 未安装")

        f = upload_root / "test.docx"
        doc = Document()
        doc.add_paragraph("Hello from docx")
        doc.save(f)

        result = await process_document(str(f))
        assert "Hello from docx" in result


class TestProcessXlsx:
    """XLSX 解析测试。"""

    @pytest.mark.asyncio
    async def test_process_xlsx(self, upload_root):
        try:
            from openpyxl import Workbook
        except ImportError:
            pytest.skip("openpyxl 未安装")

        f = upload_root / "test.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        ws.append(["a", "b", "c"])
        wb.save(f)

        result = await process_document(str(f))
        assert "Sheet1" in result
        assert "a | b | c" in result


class TestProcessPptx:
    """PPTX 解析测试。"""

    @pytest.mark.asyncio
    async def test_process_pptx(self, upload_root):
        try:
            from pptx import Presentation
        except ImportError:
            pytest.skip("python-pptx 未安装")

        f = upload_root / "test.pptx"
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "Hello from pptx"
        prs.save(f)

        result = await process_document(str(f))
        assert "Hello from pptx" in result


class TestProcessErrors:
    """异常处理测试。"""

    @pytest.mark.asyncio
    async def test_legacy_format_returns_message(self, upload_root):
        f = upload_root / "old.doc"
        f.write_text("legacy", encoding="utf-8")
        result = await process_document(str(f))
        assert "旧版格式" in result

    @pytest.mark.asyncio
    async def test_unsupported_extension_returns_empty(self, upload_root):
        f = upload_root / "unknown.xyz"
        f.write_text("data", encoding="utf-8")
        result = await process_document(str(f))
        assert result == ""

    @pytest.mark.asyncio
    async def test_nonexistent_file_raises(self, upload_root):
        missing = upload_root / "missing.txt"
        # validate_path 将 FileNotFoundError 转为 PathSecurityError，
        # process_document 再包装为 ValueError
        with pytest.raises(ValueError, match="文件路径不合法或过大"):
            await process_document(str(missing))

    @pytest.mark.asyncio
    async def test_path_outside_root_raises(self, upload_root):
        outside = upload_root / ".." / "outside.txt"
        with pytest.raises(ValueError, match="文件路径不合法或过大"):
            await process_document(str(outside))
