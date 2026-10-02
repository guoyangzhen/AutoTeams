import asyncio
import os
import logging
from typing import Optional

# P1-05-A: PyPDF2 已 EOL，迁移到 pypdf（API 兼容：PdfReader）
from pypdf import PdfReader
from pypdf.errors import PdfReadError
from docx import Document as DocxDocument
from docx.opc.exceptions import PackageNotFoundError as DocxPackageNotFoundError
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException as OpenpyxlInvalidFileException
from pptx import Presentation
from pptx.exc import PythonPptxError

from app.config import settings
from app.services.path_security import validate_path, validate_file_size, PathSecurityError
from app.utils.metrics import errors_total

logger = logging.getLogger(__name__)

LEGACY_FORMATS = {".doc", ".xls", ".ppt"}


async def process_document(path: str, allowed_root: Optional[str] = None) -> str:
    # P0-04: 路径安全校验，防止路径遍历
    # P0-07: 文件大小校验，防止 OOM
    # allowed_root: 指定允许的根目录（None 时默认 UPLOAD_ROOT）；
    #   服务端控制的示例数据目录可显式传入，绕过 UPLOAD_ROOT 限制
    try:
        safe_path = validate_path(path, allowed_root=allowed_root, must_exist=True)
        validate_file_size(safe_path, settings.MAX_FILE_SIZE_DOCUMENT)
    except PathSecurityError as e:
        logger.warning(f"文档路径校验失败: {e}")
        raise ValueError(f"文件路径不合法或过大: {e}") from e
    except FileNotFoundError:
        raise FileNotFoundError(f"文件不存在: {path}") from None
    path = safe_path

    ext = os.path.splitext(path)[1].lower()

    if ext in LEGACY_FORMATS:
        logger.warning(f"旧版Office格式 {ext} 不直接支持，请转换为新格式: {path}")
        return f"[旧版格式] 文件 {os.path.basename(path)} 使用了旧版Office格式({ext})，请转换为新格式(.docx/.xlsx/.pptx)后重新处理。"

    handlers = {
        ".pdf": _extract_pdf,
        ".docx": _extract_docx,
        ".xlsx": _extract_xlsx,
        ".pptx": _extract_pptx,
        ".txt": _extract_text,
        ".md": _extract_text,
        ".csv": _extract_text,
        ".json": _extract_text,
        ".xml": _extract_text,
        ".yaml": _extract_text,
        ".yml": _extract_text,
    }

    handler = handlers.get(ext)
    if handler is None:
        logger.warning(f"不支持的文档格式: {ext}")
        return ""

    try:
        # P2-1: 文档解析是 CPU 密集型同步操作，放到线程池避免阻塞事件循环
        return await asyncio.to_thread(handler, path)
    except (PdfReadError, DocxPackageNotFoundError, OpenpyxlInvalidFileException, PythonPptxError) as e:
        logger.error(f"文档格式损坏或无法解析 {path}: {e}", exc_info=True)
        return ""
    except (OSError, LookupError, UnicodeDecodeError) as e:
        logger.error(f"读取文档失败 {path}: {e}", exc_info=True)
        return ""
    except (RuntimeError, TypeError, KeyError) as e:
        errors_total.labels(module=__name__, exception_type=type(e).__name__).inc()
        logger.error(f"处理文档失败（未预期错误） {path}: {e}", exc_info=True)
        raise RuntimeError(f"处理文档失败: {path}") from e


def _extract_pdf(path: str) -> str:
    reader = PdfReader(path)
    pages = []
    for page in reader.pages:
        text = page.extract_text()
        if text:
            pages.append(text)
    return "\n\n".join(pages)


def _extract_docx(path: str) -> str:
    doc = DocxDocument(path)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    return "\n\n".join(paragraphs)


def _extract_xlsx(path: str) -> str:
    wb = load_workbook(path, read_only=True, data_only=True)
    all_text = []
    for sheet in wb.worksheets:
        sheet_text = [f"[{sheet.title}]"]
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                sheet_text.append(" | ".join(cells))
        if len(sheet_text) > 1:
            all_text.append("\n".join(sheet_text))
    wb.close()
    return "\n\n".join(all_text)


def _extract_pptx(path: str) -> str:
    prs = Presentation(path)
    slides_text = []
    for idx, slide in enumerate(prs.slides, 1):
        slide_text = [f"[Slide {idx}]"]
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = paragraph.text.strip()
                    if text:
                        slide_text.append(text)
        if len(slide_text) > 1:
            slides_text.append("\n".join(slide_text))
    return "\n\n".join(slides_text)


def _extract_text(path: str) -> str:
    import chardet
    with open(path, "rb") as f:
        raw = f.read()
    detected = chardet.detect(raw)
    encoding = detected.get("encoding", "utf-8") or "utf-8"
    return raw.decode(encoding, errors="replace")
